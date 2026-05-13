from __future__ import annotations

import json
import math
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
import requests
from openpyxl import load_workbook


REPO = Path(__file__).resolve().parent
FULL_WORK = REPO / "full work"
ORIGINAL_LOCAL_REPO = Path(r"C:\Users\Fiifi\Desktop\cape_verde_load_project")

LATITUDE = 14.9167
LONGITUDE = -23.5167
YEAR = 2023
EXPECTED_HOURS = 8760
OPEN_METEO_URL = (
    "https://archive-api.open-meteo.com/v1/archive"
    f"?latitude={LATITUDE}&longitude={LONGITUDE}"
    "&start_date=2023-01-01&end_date=2023-12-31"
    "&hourly=wind_speed_100m,shortwave_radiation"
    "&wind_speed_unit=ms&timezone=Atlantic%2FCape_Verde"
)
NASA_POWER_URL = (
    "https://power.larc.nasa.gov/api/temporal/hourly/point"
    f"?parameters=WS10M,ALLSKY_SFC_SW_DWN&community=RE&longitude={LONGITUDE}"
    f"&latitude={LATITUDE}&start=20230101&end=20231231&format=JSON&time-standard=LST"
)
ELECTRA_REPORT_URL = (
    "https://www.mf.gov.cv/documents/20126/4493539/"
    "RContas_Electra%2BSA%2B2023.pdf/49cabfe7-ec19-fa6f-dcea-7115b5dade7f"
    "?download=true&t=1716468826573&version=1.0"
)


@dataclass
class BuildSummary:
    load_rows: int
    weather_rows: int
    annual_load_gwh: float
    average_wind_speed: float
    annual_solar_kwh_m2: float
    nasa_succeeded: bool
    nasa_note: str
    local_workbook_created: bool
    local_workbook_path: str


def ensure_dirs(paths: Iterable[Path]) -> None:
    for path in paths:
        path.mkdir(parents=True, exist_ok=True)


def mirror_file(src: Path, relative: Path) -> None:
    destinations = [REPO / relative, FULL_WORK / relative]
    for dest in destinations:
        dest.parent.mkdir(parents=True, exist_ok=True)
        if src.resolve() == dest.resolve():
            continue
        shutil.copy2(src, dest)


def write_text_both(relative: Path, text: str) -> None:
    for base in (REPO, FULL_WORK):
        dest = base / relative
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(text, encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def find_first_existing(candidates: Iterable[Path]) -> Path | None:
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def find_load_csv() -> Path:
    candidates = [
        FULL_WORK / "santiago_hourly_load_2023_reconstructed.csv",
        REPO / "santiago_hourly_load_2023_reconstructed.csv",
        ORIGINAL_LOCAL_REPO / "full work" / "santiago_hourly_load_2023_reconstructed.csv",
        Path(r"C:\Users\Fiifi\Downloads\CapeVerdeReferenceSystemData_v002\Santiago_Load_Output\santiago_hourly_load_2023_reconstructed.csv"),
    ]
    found = find_first_existing(candidates)
    if not found:
        raise FileNotFoundError("Existing Santiago reconstructed load CSV was not found. Load was not regenerated.")
    return found


def find_load_xlsx() -> Path | None:
    candidates = [
        FULL_WORK / "santiago_hourly_load_2023_reconstructed.xlsx",
        REPO / "santiago_hourly_load_2023_reconstructed.xlsx",
        ORIGINAL_LOCAL_REPO / "full work" / "santiago_hourly_load_2023_reconstructed.xlsx",
        Path(r"C:\Users\Fiifi\Downloads\CapeVerdeReferenceSystemData_v002\Santiago_Load_Output\santiago_hourly_load_2023_reconstructed.xlsx"),
    ]
    return find_first_existing(candidates)


def copy_existing_load_evidence() -> pd.DataFrame:
    load_csv = find_load_csv()
    load_xlsx = find_load_xlsx()
    mirror_file(load_csv, Path("data/processed/model_inputs/santiago_hourly_load_2023_reconstructed.csv"))
    if load_xlsx:
        mirror_file(load_xlsx, Path("data/processed/model_inputs/santiago_hourly_load_2023_reconstructed.xlsx"))

    load = pd.read_csv(load_csv)
    if len(load) != EXPECTED_HOURS:
        raise ValueError(f"Load file has {len(load)} rows, expected {EXPECTED_HOURS}.")
    load["timestamp"] = pd.to_datetime(load["timestamp"])
    if "energy_mwh" in load.columns:
        load_mwh = load["energy_mwh"].astype(float)
    elif "load_mw" in load.columns:
        load_mwh = load["load_mw"].astype(float)
    else:
        raise ValueError("Load file does not have energy_mwh or load_mw.")
    annual_gwh = load_mwh.sum() / 1000
    if not (263.7 <= annual_gwh <= 264.0):
        raise ValueError(f"Annual load is {annual_gwh:.6f} GWh, not close to 263.839 GWh.")
    if load["timestamp"].duplicated().any():
        raise ValueError("Load file has duplicate timestamps.")
    if load["timestamp"].sort_values().tolist() != load["timestamp"].tolist():
        raise ValueError("Load timestamps are not sorted.")
    return load


def copy_raw_load_sources() -> None:
    downloads_root = Path(r"C:\Users\Fiifi\Downloads")
    dtu_xlsx = find_first_existing(
        [
            Path(r"C:\Users\Fiifi\Downloads\CapeVerdeReferenceSystemData_v002\CapeVerdeReferenceSystemData_v002\SantiagoData_v4.xlsx"),
            ORIGINAL_LOCAL_REPO / "CapeVerdeReferenceSystemData_v002" / "SantiagoData_v4.xlsx",
        ]
    )
    if not dtu_xlsx:
        matches = list(downloads_root.rglob("SantiagoData_v4.xlsx"))
        dtu_xlsx = matches[0] if matches else None
    if dtu_xlsx:
        mirror_file(dtu_xlsx, Path("data/raw/load/SantiagoData_v4.xlsx"))

    dtu_zip = find_first_existing(
        [
            Path(r"C:\Users\Fiifi\Downloads\CapeVerdeReferenceSystemData_v002.zip"),
            ORIGINAL_LOCAL_REPO / "CapeVerdeReferenceSystemData_v002.zip",
        ]
    )
    if dtu_zip:
        mirror_file(dtu_zip, Path("data/raw/load/CapeVerdeReferenceSystemData_v002.zip"))

    note = """# Load Source Note

Load shape source: DTU Cape Verde Reference System.

Used file: `SantiagoData_v4.xlsx`.

Used sheet: `Equivalent Week Profile Load`.

Scaling source: Electra 2023 Santiago annual production.

Important note: this is a reconstructed Santiago load curve, not official SCADA hourly data. SCADA means the operator's real control-room metering and dispatch record.
"""
    write_text_both(Path("data/raw/load/load_source_note.md"), note)


def download_file(url: str, dest: Path, timeout: int = 120) -> str:
    dest.parent.mkdir(parents=True, exist_ok=True)
    response = requests.get(url, timeout=timeout)
    response.raise_for_status()
    dest.write_bytes(response.content)
    return response.headers.get("content-type", "")


def download_open_meteo() -> pd.DataFrame:
    raw_json = REPO / "data/raw/wind_solar/open_meteo_santiago_2023_hourly_weather_raw.json"
    raw_json.parent.mkdir(parents=True, exist_ok=True)
    response = requests.get(OPEN_METEO_URL, timeout=120)
    response.raise_for_status()
    payload = response.json()
    raw_json.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    write_text(REPO / "data/raw/wind_solar/open_meteo_request_url.txt", OPEN_METEO_URL + "\n")

    units = payload.get("hourly_units", {})
    wind_unit = units.get("wind_speed_100m", "")
    if wind_unit not in {"m/s", "ms"}:
        raise ValueError(f"Open-Meteo wind_speed_100m unit is {wind_unit!r}; URL must use wind_speed_unit=ms.")

    hourly = payload.get("hourly", {})
    required = ["time", "wind_speed_100m", "shortwave_radiation"]
    missing = [key for key in required if key not in hourly]
    if missing:
        raise ValueError(f"Open-Meteo response missing fields: {missing}")

    weather = pd.DataFrame(
        {
            "timestamp": pd.to_datetime(hourly["time"]),
            "wind_speed_100m_m_s": pd.to_numeric(hourly["wind_speed_100m"], errors="raise"),
            "solar_radiation_w_m2": pd.to_numeric(hourly["shortwave_radiation"], errors="raise"),
        }
    )
    weather.insert(0, "hour_of_year", range(1, len(weather) + 1))
    weather["solar_radiation_kw_m2"] = weather["solar_radiation_w_m2"] / 1000.0
    weather["source"] = "Open-Meteo Historical Weather API, wind_speed_100m and shortwave_radiation, Praia/Santiago coordinates"

    validate_weather(weather)
    weather_path = REPO / "data/processed/model_inputs/santiago_wind_solar_2023_8760.csv"
    weather_path.parent.mkdir(parents=True, exist_ok=True)
    weather.to_csv(weather_path, index=False)

    # Mirror raw and processed files into full work.
    for relative in [
        Path("data/raw/wind_solar/open_meteo_santiago_2023_hourly_weather_raw.json"),
        Path("data/raw/wind_solar/open_meteo_request_url.txt"),
        Path("data/processed/model_inputs/santiago_wind_solar_2023_8760.csv"),
    ]:
        src = REPO / relative
        dest = FULL_WORK / relative
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dest)
    return weather


def validate_weather(weather: pd.DataFrame) -> None:
    if len(weather) != EXPECTED_HOURS:
        raise ValueError(f"Weather has {len(weather)} rows, expected {EXPECTED_HOURS}.")
    if weather["timestamp"].duplicated().any():
        raise ValueError("Weather has duplicate timestamps.")
    expected = pd.date_range("2023-01-01 00:00", "2023-12-31 23:00", freq="h")
    if not weather["timestamp"].equals(pd.Series(expected)):
        missing = sorted(set(expected) - set(weather["timestamp"]))
        raise ValueError(f"Weather timestamps are not exactly the 2023 hourly series. Missing count: {len(missing)}")
    if weather[["wind_speed_100m_m_s", "solar_radiation_w_m2", "solar_radiation_kw_m2"]].isna().any().any():
        raise ValueError("Weather data has missing numeric values.")
    if (weather["solar_radiation_w_m2"] < 0).any():
        raise ValueError("Weather data has negative solar radiation values.")
    if (weather["wind_speed_100m_m_s"] < 0).any() or (weather["wind_speed_100m_m_s"] > 80).any():
        raise ValueError("Weather data has impossible wind-speed values.")
    check = weather["solar_radiation_w_m2"] / 1000.0
    if not (check.round(12).equals(weather["solar_radiation_kw_m2"].round(12))):
        raise ValueError("Solar conversion check failed.")


def make_combined_inputs(load: pd.DataFrame, weather: pd.DataFrame) -> pd.DataFrame:
    load2 = load.copy()
    load2["timestamp"] = pd.to_datetime(load2["timestamp"])
    load2["load_mwh"] = (
        pd.to_numeric(load2["energy_mwh"], errors="raise")
        if "energy_mwh" in load2.columns
        else pd.to_numeric(load2["load_mw"], errors="raise")
    )
    combined = pd.DataFrame(
        {
            "hour_of_year": range(1, EXPECTED_HOURS + 1),
            "timestamp": weather["timestamp"],
            "load_mwh": load2["load_mwh"].to_numpy(),
            "wind_speed_m_s": weather["wind_speed_100m_m_s"].to_numpy(),
            "solar_radiation_kw_m2": weather["solar_radiation_kw_m2"].to_numpy(),
        }
    )
    if not load2["timestamp"].equals(weather["timestamp"]):
        raise ValueError("Load and weather timestamps do not match exactly.")
    validate_combined(combined)
    out = REPO / "data/processed/model_inputs/re100_santiago_combined_inputs_2023.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    combined.to_csv(out, index=False)
    mirror_file(out, Path("data/processed/model_inputs/re100_santiago_combined_inputs_2023.csv"))
    return combined


def validate_combined(combined: pd.DataFrame) -> None:
    if len(combined) != EXPECTED_HOURS:
        raise ValueError(f"Combined inputs have {len(combined)} rows, expected {EXPECTED_HOURS}.")
    annual_mwh = combined["load_mwh"].sum()
    if not (263700 <= annual_mwh <= 264000):
        raise ValueError(f"Combined annual load is {annual_mwh:.3f} MWh, not close to 263,839.489 MWh.")
    if combined[["load_mwh", "wind_speed_m_s", "solar_radiation_kw_m2"]].isna().any().any():
        raise ValueError("Combined inputs have missing numeric values.")
    if (combined["wind_speed_m_s"] < 0).any() or (combined["wind_speed_m_s"] > 80).any():
        raise ValueError("Combined inputs have impossible wind speed.")
    if (combined["solar_radiation_kw_m2"] < 0).any():
        raise ValueError("Combined inputs have negative solar radiation.")
    expected = pd.date_range("2023-01-01 00:00", "2023-12-31 23:00", freq="h")
    if not pd.to_datetime(combined["timestamp"]).equals(pd.Series(expected)):
        raise ValueError("Combined timestamp order is not the exact 2023 hourly series.")


def write_mapping_note() -> None:
    mapping = """# RE100 Input Mapping for Santiago 2023

## Load

Goes into: `Hourly Load!F7:F8766`

Unit: MWh/h, same as MW for one hour.

Source: reconstructed Santiago 2023 hourly load curve from DTU Santiago profile scaled to Electra's official Santiago annual production.

## Wind

Goes into: `Wind and Solar Input!D5:D8764`

Unit: m/s.

Source variable: Open-Meteo `wind_speed_100m`.

Important: set the measurement height in Prof. Hohmeyer's model to 100 m if using `wind_speed_100m`. The matching input cell is `Input & Output!C66 = 100`.

## Solar

Goes into: `Wind and Solar Input!E5:E8764`

Unit: kW/m2.

Source variable: Open-Meteo `shortwave_radiation`.

Conversion: Open-Meteo gives W/m2, so the model input is W/m2 divided by 1000.

## Caution

These weather data are Open-Meteo reanalysis/modelled historical weather data, not measured wind mast or on-site pyranometer observations. They are usable public hourly inputs, but they are not verified site measurements.
"""
    write_text_both(Path("data/notes/re100_input_mapping.md"), mapping)


def attempt_nasa_cross_check(open_meteo_weather: pd.DataFrame) -> tuple[bool, str]:
    raw_path = REPO / "data/raw/wind_solar/nasa_power_santiago_2023_hourly_raw.json"
    comparison_path = REPO / "data/processed/validation/nasa_power_open_meteo_comparison.csv"
    try:
        response = requests.get(NASA_POWER_URL, timeout=120)
        response.raise_for_status()
        payload = response.json()
        raw_path.parent.mkdir(parents=True, exist_ok=True)
        raw_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        params = payload["properties"]["parameter"]
        ws10m = pd.Series(params["WS10M"], dtype="float64")
        solar = pd.Series(params["ALLSKY_SFC_SW_DWN"], dtype="float64")
        if len(ws10m) != EXPECTED_HOURS or len(solar) != EXPECTED_HOURS:
            raise ValueError(f"NASA returned WS10M={len(ws10m)} rows and solar={len(solar)} rows.")

        nasa_timestamps = pd.to_datetime(ws10m.index, format="%Y%m%d%H")
        nasa = pd.DataFrame(
            {
                "timestamp": nasa_timestamps,
                "nasa_ws10m_m_s": ws10m.to_numpy(),
                "nasa_solar_raw": solar.to_numpy(),
            }
        )
        om = open_meteo_weather[["timestamp", "wind_speed_100m_m_s", "solar_radiation_kw_m2"]].copy()
        merged = om.merge(nasa, on="timestamp", how="inner")
        if len(merged) != EXPECTED_HOURS:
            raise ValueError(f"NASA/Open-Meteo merge has {len(merged)} rows.")

        # NASA POWER hourly ALLSKY_SFC_SW_DWN is treated as Wh/m2 per hour here.
        merged["open_meteo_solar_kwh_m2"] = merged["solar_radiation_kw_m2"]
        merged["nasa_solar_kwh_m2"] = merged["nasa_solar_raw"] / 1000.0
        merged["month"] = merged["timestamp"].dt.month

        rows: list[dict[str, object]] = [
            {
                "metric": "annual_average_wind_speed",
                "period": "2023",
                "open_meteo_value": merged["wind_speed_100m_m_s"].mean(),
                "nasa_power_value": merged["nasa_ws10m_m_s"].mean(),
                "unit": "m/s",
                "notes": "Open-Meteo is 100 m wind speed. NASA POWER value is WS10M at 10 m, so this is a broad cross-check only.",
            },
            {
                "metric": "annual_solar_radiation_sum",
                "period": "2023",
                "open_meteo_value": merged["open_meteo_solar_kwh_m2"].sum(),
                "nasa_power_value": merged["nasa_solar_kwh_m2"].sum(),
                "unit": "kWh/m2 per year",
                "notes": "Open-Meteo W/m2 converted to kWh/m2 for hourly sums. NASA POWER ALLSKY_SFC_SW_DWN treated as Wh/m2 per hour.",
            },
        ]
        monthly = merged.groupby("month", as_index=False)[["open_meteo_solar_kwh_m2", "nasa_solar_kwh_m2"]].sum()
        for _, row in monthly.iterrows():
            rows.append(
                {
                    "metric": "monthly_solar_radiation_sum",
                    "period": int(row["month"]),
                    "open_meteo_value": row["open_meteo_solar_kwh_m2"],
                    "nasa_power_value": row["nasa_solar_kwh_m2"],
                    "unit": "kWh/m2 per month",
                    "notes": "Monthly solar total cross-check. Different datasets and definitions can differ.",
                }
            )
        comparison = pd.DataFrame(rows)
        comparison_path.parent.mkdir(parents=True, exist_ok=True)
        comparison.to_csv(comparison_path, index=False)
        for relative in [
            Path("data/raw/wind_solar/nasa_power_santiago_2023_hourly_raw.json"),
            Path("data/processed/validation/nasa_power_open_meteo_comparison.csv"),
        ]:
            src = REPO / relative
            dest = FULL_WORK / relative
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest)
        note = f"NASA POWER cross-check succeeded. Request URL:\n\n{NASA_POWER_URL}\n"
        write_text_both(Path("data/raw/wind_solar/nasa_power_cross_check_note.md"), note)
        return True, note
    except Exception as exc:
        note = f"NASA POWER cross-check was attempted but not completed.\n\nExact request URL:\n{NASA_POWER_URL}\n\nExact error:\n{exc}\n"
        write_text_both(Path("data/raw/wind_solar/nasa_power_cross_check_note.md"), note)
        return False, str(exc)


def download_electra_report() -> None:
    report_dest = REPO / "data/raw/reports/electra_2023_annual_report.pdf"
    try:
        download_file(ELECTRA_REPORT_URL, report_dest)
        mirror_file(report_dest, Path("data/raw/reports/electra_2023_annual_report.pdf"))
    except Exception as exc:
        note = f"Electra report download was attempted but not completed.\n\nURL:\n{ELECTRA_REPORT_URL}\n\nError:\n{exc}\n"
        write_text_both(Path("data/raw/reports/electra_report_download_note.md"), note)


def write_electricity_maps_note() -> None:
    note = """# Electricity Maps Note

Electricity Maps Cabo Verde data is only a benchmark source for the national Cabo Verde zone. It is not final Santiago hourly load data and it is not official Santiago operator SCADA.

The final RE100 Santiago input file uses the Santiago-specific DTU reconstructed load curve plus Open-Meteo hourly wind and solar data.
"""
    write_text_both(Path("data/raw/electricity_maps/electricity_maps_note.md"), note)


def fill_local_professor_workbook(combined: pd.DataFrame) -> tuple[bool, str]:
    candidates = [
        ORIGINAL_LOCAL_REPO / "Model RE100 and handbook" / "RE100 Model Version 1.0 2026.xlsx",
        REPO / "Model RE100 and handbook" / "RE100 Model Version 1.0 2026.xlsx",
        Path(r"C:\Users\Fiifi\Downloads\CapeVerdeReferenceSystemData_v002\CapeVerdeReferenceSystemData_v002\RE100 Model Version 1.0 2026.xlsx"),
    ]
    source = find_first_existing(candidates)
    if not source:
        return False, "Professor workbook not found locally."

    output = ORIGINAL_LOCAL_REPO / "RE100_Model_Santiago_2023_filled_LOCAL_ONLY.xlsx"
    shutil.copy2(source, output)
    wb = load_workbook(output)
    required_sheets = ["Hourly Load", "Wind and Solar Input", "Input & Output"]
    missing = [sheet for sheet in required_sheets if sheet not in wb.sheetnames]
    if missing:
        raise ValueError(f"Professor workbook copy missing required sheets: {missing}")

    load_ws = wb["Hourly Load"]
    weather_ws = wb["Wind and Solar Input"]
    io_ws = wb["Input & Output"]
    for idx, value in enumerate(combined["load_mwh"], start=7):
        load_ws[f"F{idx}"] = float(value)
    for idx, value in enumerate(combined["wind_speed_m_s"], start=5):
        weather_ws[f"D{idx}"] = float(value)
    for idx, value in enumerate(combined["solar_radiation_kw_m2"], start=5):
        weather_ws[f"E{idx}"] = float(value)
    io_ws["C66"] = 100
    wb.save(output)
    return True, str(output)


def write_source_links() -> None:
    source_links = f"""# Source Links

## 1. DTU Cape Verde Reference System

Use: Santiago load profile.

File: `SantiagoData_v4.xlsx`.

Sheet: `Equivalent Week Profile Load`.

Link: https://data.dtu.dk/articles/dataset/Cape_Verde_Reference_System_v2_0/17430413

## 2. Electra 2023 Annual Report

Use: Santiago 2023 annual production and peak validation.

Page: page 25 for annual Santiago production.

Page: page 26 for Santiago peak.

Link: {ELECTRA_REPORT_URL}

## 3. Prof. Hohmeyer RE100 Handbook

Use: tells us where model inputs go.

Pages: page 15 and page 16 for hourly load, wind speed, and solar radiation inputs.

Pages: page 19 to page 22 for economic and technical input cells.

Sharing note: the handbook is not uploaded here because permission to redistribute it was not confirmed.

## 4. Open-Meteo Historical Weather API

Use: hourly wind speed and solar radiation for Santiago 2023.

Documentation: https://open-meteo.com/en/docs/historical-weather-api

Exact request URL used:

{OPEN_METEO_URL}

## 5. NASA POWER Hourly API

Use: independent weather cross-check if downloaded.

Documentation: https://power.larc.nasa.gov/docs/services/api/temporal/hourly/

Exact request URL attempted:

{NASA_POWER_URL}

## 6. Electricity Maps Cabo Verde

Use: benchmark only, not final Santiago load source.

Important note: Electricity Maps Cabo Verde represents the national Cabo Verde zone, not Santiago-only operator data. It is not used as final Santiago load input.

Documentation: https://portal.electricitymaps.com/docs/api

Methodology: https://www.electricitymaps.com/research
"""
    write_text(REPO / "data/sources/source_links.md", source_links)
    write_text(FULL_WORK / "sources/source_links.md", source_links)
    write_text(FULL_WORK / "data/sources/source_links.md", source_links)


def update_readme() -> None:
    readme = FULL_WORK / "README_FULL_WORK.md"
    if not readme.exists():
        return
    text = readme.read_text(encoding="utf-8")
    section_title = "## Added Wind and Solar Data for Santiago"
    section = f"""

{section_title}

We added 2023 hourly wind and solar data for Santiago using the Open-Meteo Historical Weather API.

- Year: 2023, with 8760 hourly rows.
- Location: Praia/Santiago default coordinates, latitude {LATITUDE}, longitude {LONGITUDE}.
- Wind: `wind_speed_100m` in m/s.
- Solar: `shortwave_radiation` from Open-Meteo in W/m2, converted to kW/m2 by dividing by 1000.
- Processed weather file: `data/processed/model_inputs/santiago_wind_solar_2023_8760.csv`.
- Combined RE100 input file: `data/processed/model_inputs/re100_santiago_combined_inputs_2023.csv`.
- Raw Open-Meteo file: `data/raw/wind_solar/open_meteo_santiago_2023_hourly_weather_raw.json`.
- Exact Open-Meteo request URL: `{OPEN_METEO_URL}`.
- The professor workbook is not uploaded unless sharing permission is confirmed. A local-only filled copy may exist on this laptop.
"""
    if section_title in text:
        text = text[: text.index(section_title)].rstrip() + section
    else:
        text = text.rstrip() + section
    readme.write_text(text + "\n", encoding="utf-8")


def write_human_methodology() -> None:
    text = f"""# Methodology for Santiago Wind and Solar RE100 Inputs

## 1. What we needed

We needed the other hourly data that Prof. Hohmeyer's RE100 model needs for Santiago. Thiss means hourly wind speed and hourly solar radiation for the same 8760-hour year as the load curve.

## 2. How we got load

We did not rebuild the load curve. We used the existing Santiago reconstructed 2023 load file already in the project. It has 8760 rows and an annual total close to 263.839 GWh.

## 3. How we got wind

We used the Open-Meteo Historical Weather API for Santiago/Praia coordinates. The wind variable is `wind_speed_100m`, which means wind speed 100 m above ground. This is modelled or reanalysis weather data, not a measured wind mast file.

## 4. How we got solar

We used Open-Meteo `shortwave_radiation`. Solar radiation means sunlight energy reaching the ground. This is also modelled or reanalysis weather data, not a verified solar station measurement.

## 5. How we converted the units

Open-Meteo gave solar radiation in W/m2. Prof. Hohmeyer's model needs kW/m2, so we divided by 1000. The wind was requested directly in m/s, so no speed conversion was needed.

## 6. Where the values go in Prof. Hohmeyer's model

Load goes into `Hourly Load!F7:F8766`.

Wind goes into `Wind and Solar Input!D5:D8764`.

Solar goes into `Wind and Solar Input!E5:E8764`.

Because the wind source is 100 m wind speed, the model measurement height should be set to 100 m in `Input & Output!C66`.

## 7. What is official

The official load anchors are Electra's Santiago 2023 annual production and peak values. Those come from the Electra 2023 annual report.

## 8. What is reconstructed

The hourly load shape is reconstructed from DTU Santiago-specific load-profile data. The wind and solar are public hourly weather model or reanalysis data from Open-Meteo. They are not official measured Santiago wind and solar station data.

## 9. What still needs care

We should not call this exact measured 8760-hour Santiago energy-system data. It is a strong public input pack for running the professor's model, but official operator SCADA and measured site weather data would be better if they become available. A small NASA POWER cross-check was also attempted so the weather input is not just accepted without any outside comparison.
"""
    write_text(FULL_WORK / "methodology_RE100_Santiago_other_data.md", text)


def write_prompt_file() -> None:
    text = """Prompt summary: add Santiago wind and solar data for Prof. Hohmeyer's RE100 model.

Main goals:
- Do not build a new model.
- Keep the existing Santiago hourly load file if it is valid.
- Download 2023 hourly wind_speed_100m and shortwave_radiation from Open-Meteo for Praia/Santiago.
- Convert shortwave_radiation from W/m2 to kW/m2.
- Build 8760-row wind/solar and combined RE100 input CSV files.
- Save raw source files and exact request URLs.
- Try NASA POWER as a weather cross-check.
- Fill a local-only copy of the professor workbook if present, but do not commit it.
- Add clear source notes, a simple presentation, and GitHub evidence files.
"""
    write_text(FULL_WORK / "prompts/codex_prompt_RE100_Santiago.txt", text)


def create_simple_presentation(combined: pd.DataFrame) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    out = FULL_WORK / "presentation/Santiago_RE100_Data_Added_Simple_Explanation.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_color = RGBColor(0, 0, 0)
    body_color = RGBColor(40, 40, 40)

    def add_slide(title: str, bullets: list[str], source: str | None = None) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.65))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = title_color

        bb = slide.shapes.add_textbox(Inches(0.95), Inches(1.35), Inches(11.6), Inches(4.95))
        tf = bb.text_frame
        tf.clear()
        for i, bullet in enumerate(bullets):
            bp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            bp.text = bullet
            bp.font.size = Pt(21)
            bp.font.color.rgb = body_color
            bp.space_after = Pt(10)
        if source:
            sb = slide.shapes.add_textbox(Inches(0.75), Inches(6.65), Inches(12), Inches(0.45))
            sp = sb.text_frame.paragraphs[0]
            sp.text = "Source: " + source
            sp.font.size = Pt(11)
            sp.font.color.rgb = RGBColor(90, 90, 90)

    add_slide(
        "How we prepared Santiago data for Prof. Hohmeyer's RE100 model",
        [
            "We did not build a new model.",
            "We used Prof. Hohmeyer's RE100 Excel model.",
            "We replaced the Barbados data with Santiago, Cabo Verde data.",
        ],
    )
    add_slide(
        "What the model needs",
        [
            "Hourly load: how much electricity Santiago needs each hour.",
            "Hourly wind speed: how much wind is available each hour.",
            "Hourly solar radiation: how much sunlight is available each hour.",
            "One normal year has 8760 hours.",
        ],
        "Prof. Hohmeyer RE100 Handbook, pages 15 and 16.",
    )
    add_slide(
        "The load data we already prepared",
        [
            "We used the DTU Santiago load profile.",
            "DTU means Technical University of Denmark.",
            "We scaled it to Electra's official 2023 Santiago annual production.",
            "We checked the peak against Electra's 2023 Santiago peak.",
        ],
        "DTU dataset: https://data.dtu.dk/articles/dataset/Cape_Verde_Reference_System_v2_0/17430413. Electra report pages 25 and 26.",
    )
    add_slide(
        "Why we added wind data",
        [
            "The model needs wind speed for every hour.",
            "We used Open-Meteo Historical Weather API.",
            "We used Santiago coordinates.",
            "We downloaded 2023 hourly wind speed at 100 m.",
            "Because this is 100 m wind, the model input height must be 100 m.",
        ],
        f"Open-Meteo docs: https://open-meteo.com/en/docs/historical-weather-api. Request: {OPEN_METEO_URL}",
    )
    add_slide(
        "Why we added solar data",
        [
            "The model needs solar radiation for every hour.",
            "Solar radiation means sunlight energy reaching the ground.",
            "Open-Meteo gives it in W/m2.",
            "The model needs kW/m2.",
            "So we divided by 1000.",
        ],
        f"Open-Meteo docs: https://open-meteo.com/en/docs/historical-weather-api. Request: {OPEN_METEO_URL}",
    )
    add_slide(
        "Where the data goes in the model",
        [
            "Load goes into Hourly Load column F.",
            "Wind goes into Wind and Solar Input column D.",
            "Solar goes into Wind and Solar Input column E.",
            "We kept the professor's formulas unchanged.",
        ],
        "Prof. Hohmeyer RE100 Handbook, pages 15 and 16.",
    )
    add_slide(
        "What is official and what is not",
        [
            "Official: Electra annual Santiago 2023 production.",
            "Official: Electra Santiago 2023 peak.",
            "Research source: DTU Santiago load profile.",
            "Weather source: Open-Meteo hourly wind and solar.",
            "Not official: exact 8760-hour SCADA load data, because it was not public.",
            "SCADA means the operator's real control-room meter data.",
        ],
    )
    add_slide(
        "What we put in GitHub",
        [
            "raw data files",
            "processed 8760-hour model input files",
            "source links",
            "method notes",
            "presentation",
            "professor workbook not uploaded unless permission is confirmed",
        ],
    )
    add_slide(
        "What we can defend",
        [
            "We can defend the annual load total.",
            "We can defend the peak check.",
            "We can defend where wind and solar came from.",
            "We can defend how units were converted.",
            "We cannot call reconstructed hourly load official meter data without Electra's real file.",
        ],
    )
    add_slide(
        "Final simple conclusion",
        [
            "We prepared the Santiago input data needed for the RE100 model.",
            "The model can now use Santiago load, wind, and solar instead of Barbados data.",
            "The work is documented so every value can be traced to a source.",
        ],
    )
    prs.save(out)
    return out


def write_validation_summary(load: pd.DataFrame, weather: pd.DataFrame, combined: pd.DataFrame, nasa_succeeded: bool, nasa_note: str) -> None:
    rows = [
        {"check": "load_rows", "value": len(load), "expected": EXPECTED_HOURS, "status": "pass" if len(load) == EXPECTED_HOURS else "fail"},
        {"check": "wind_solar_rows", "value": len(weather), "expected": EXPECTED_HOURS, "status": "pass" if len(weather) == EXPECTED_HOURS else "fail"},
        {"check": "combined_rows", "value": len(combined), "expected": EXPECTED_HOURS, "status": "pass" if len(combined) == EXPECTED_HOURS else "fail"},
        {"check": "annual_load_gwh", "value": combined["load_mwh"].sum() / 1000, "expected": "about 263.839", "status": "pass"},
        {"check": "average_wind_speed_m_s", "value": combined["wind_speed_m_s"].mean(), "expected": "not negative, plausible", "status": "pass"},
        {"check": "annual_solar_kwh_m2", "value": combined["solar_radiation_kw_m2"].sum(), "expected": "not negative", "status": "pass"},
        {"check": "nasa_power_cross_check", "value": "succeeded" if nasa_succeeded else "failed", "expected": "attempted", "status": "pass", "notes": nasa_note},
    ]
    validation = pd.DataFrame(rows)
    out = REPO / "data/processed/validation/re100_santiago_wind_solar_validation.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    validation.to_csv(out, index=False)
    mirror_file(out, Path("data/processed/validation/re100_santiago_wind_solar_validation.csv"))


def main() -> None:
    ensure_dirs(
        [
            REPO / "data/raw/load",
            REPO / "data/raw/wind_solar",
            REPO / "data/raw/reports",
            REPO / "data/raw/electricity_maps",
            REPO / "data/processed/model_inputs",
            REPO / "data/processed/validation",
            REPO / "data/sources",
            REPO / "data/notes",
            FULL_WORK / "data/raw/load",
            FULL_WORK / "data/raw/wind_solar",
            FULL_WORK / "data/raw/reports",
            FULL_WORK / "data/raw/electricity_maps",
            FULL_WORK / "data/processed/model_inputs",
            FULL_WORK / "data/processed/validation",
            FULL_WORK / "data/sources",
            FULL_WORK / "data/notes",
        ]
    )
    load = copy_existing_load_evidence()
    copy_raw_load_sources()
    weather = download_open_meteo()
    combined = make_combined_inputs(load, weather)
    write_mapping_note()
    nasa_succeeded, nasa_note = attempt_nasa_cross_check(weather)
    download_electra_report()
    write_electricity_maps_note()
    local_created, local_path = fill_local_professor_workbook(combined)
    write_source_links()
    update_readme()
    write_human_methodology()
    write_prompt_file()
    create_simple_presentation(combined)
    write_validation_summary(load, weather, combined, nasa_succeeded, nasa_note)

    summary = BuildSummary(
        load_rows=len(load),
        weather_rows=len(weather),
        annual_load_gwh=combined["load_mwh"].sum() / 1000,
        average_wind_speed=combined["wind_speed_m_s"].mean(),
        annual_solar_kwh_m2=combined["solar_radiation_kw_m2"].sum(),
        nasa_succeeded=nasa_succeeded,
        nasa_note=nasa_note,
        local_workbook_created=local_created,
        local_workbook_path=local_path,
    )
    summary_path = REPO / "data/processed/validation/re100_santiago_wind_solar_build_summary.json"
    summary_path.write_text(json.dumps(summary.__dict__, indent=2), encoding="utf-8")
    mirror_file(summary_path, Path("data/processed/validation/re100_santiago_wind_solar_build_summary.json"))

    print(json.dumps(summary.__dict__, indent=2))


if __name__ == "__main__":
    main()
